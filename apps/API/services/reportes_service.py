from __future__ import annotations

import io
import logging
from collections.abc import Sequence
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

from apps.API.core.timezone import now as tz_now
from apps.API.models.asistencia import Asistencia
from apps.API.models.novedad import Novedad
from apps.API.models.reporte_semanal import ReporteSemanal
from apps.API.repositories import (
    auditoria_repository,
    reporte_repository,
)
from shared.constants.operacion_auditoria import OperacionAuditoria
from shared.constants.recurso_auditoria import RecursoAuditoria
from shared.exceptions.reportes import (
    RangoFechasInvalidoError,
    ReporteDuplicadoError,
    ReporteGeneracionError,
    ReporteNoEncontradoError,
    ReporteSinDatosError,
)

logger = logging.getLogger(__name__)

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe_cell(value: str | None) -> str:
    if value is None:
        return ""
    text = str(value)
    if text and text[0] in _FORMULA_PREFIXES:
        return "'" + text
    return text

_EXCEL_HEADERS = [
    "Documento",
    "Empleado",
    "Sede",
    "Tipo Registro",
    "Fecha",
    "Hora Registro",
    "Novedad",
]


def _validar_rango_fechas(fecha_desde: date | None, fecha_hasta: date | None) -> None:
    if fecha_desde is not None and fecha_hasta is not None and fecha_desde > fecha_hasta:
        raise RangoFechasInvalidoError()


def _lunes_de_semana(referencia: date) -> date:
    return referencia - timedelta(days=referencia.weekday())


async def consultar_asistencia(
    session: AsyncSession,
    *,
    id_empleado: int | None = None,
    id_sede: int | None = None,
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
) -> tuple[Sequence[Asistencia], dict[int, Novedad]]:
    _validar_rango_fechas(fecha_desde, fecha_hasta)

    asistencias = await reporte_repository.get_asistencias_filtradas(
        session,
        id_empleado=id_empleado,
        id_sede=id_sede,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
    )

    ids = [a.id for a in asistencias]
    novedades_map = await reporte_repository.get_novedades_por_asistencia_ids(
        session, ids
    )

    return asistencias, novedades_map


async def generar_reporte_semanal(
    session: AsyncSession,
    *,
    fecha_referencia: date | None = None,
    id_usuario: int | None = None,
) -> ReporteSemanal:
    hoy = fecha_referencia or tz_now().date()
    lunes_anterior = _lunes_de_semana(hoy) - timedelta(weeks=1)
    domingo_anterior = lunes_anterior + timedelta(days=6)

    existente = await reporte_repository.get_reporte_by_periodo(
        session, lunes_anterior, domingo_anterior
    )
    if existente is not None:
        raise ReporteDuplicadoError(
            str(lunes_anterior), str(domingo_anterior)
        )

    try:
        asistencias = await reporte_repository.get_asistencias_filtradas(
            session, fecha_desde=lunes_anterior, fecha_hasta=domingo_anterior
        )
        total_novedades = await reporte_repository.contar_novedades_periodo(
            session, lunes_anterior, domingo_anterior
        )

        timestamp = tz_now()
        reporte = await reporte_repository.create_reporte(
            session,
            fecha_inicio=lunes_anterior,
            fecha_fin=domingo_anterior,
            total_registros=len(asistencias),
            total_novedades=total_novedades,
            now=timestamp,
        )
    except ReporteDuplicadoError:
        raise
    except Exception as exc:
        logger.error("Error generando reporte semanal: %s", exc, exc_info=True)
        raise ReporteGeneracionError("Error interno al generar el reporte") from exc

    await auditoria_repository.create(
        session,
        id_usuario=id_usuario,
        recurso=RecursoAuditoria.REPORTE,
        id_recurso=str(reporte.id),
        operacion=OperacionAuditoria.INSERT,
        valor_nuevo={
            "periodo_inicio": str(lunes_anterior),
            "periodo_fin": str(domingo_anterior),
            "total_registros": reporte.total_registros,
            "total_novedades": reporte.total_novedades,
        },
        timestamp_accion=timestamp,
    )

    logger.info(
        "Reporte semanal generado: periodo=%s - %s, registros=%d, novedades=%d",
        lunes_anterior,
        domingo_anterior,
        reporte.total_registros,
        reporte.total_novedades,
    )

    return reporte


async def list_reportes_semanales(
    session: AsyncSession,
) -> Sequence[ReporteSemanal]:
    return await reporte_repository.get_all_reportes(session)


async def get_reporte_semanal(
    session: AsyncSession, reporte_id: int
) -> ReporteSemanal:
    reporte = await reporte_repository.get_reporte_by_id(session, reporte_id)
    if reporte is None:
        raise ReporteNoEncontradoError(reporte_id)
    return reporte


def generar_excel(
    asistencias: Sequence[Asistencia],
    novedades_map: dict[int, Novedad],
    *,
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
) -> bytes:
    if not asistencias:
        raise ReporteSinDatosError()

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte Asistencia"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")

    for col_idx, header in enumerate(_EXCEL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, asistencia in enumerate(asistencias, start=2):
        novedad = novedades_map.get(asistencia.id)

        ws.cell(row=row_idx, column=1, value=_safe_cell(asistencia.empleado.documento))
        ws.cell(
            row=row_idx,
            column=2,
            value=_safe_cell(f"{asistencia.empleado.nombre} {asistencia.empleado.apellido}"),
        )
        ws.cell(row=row_idx, column=3, value=_safe_cell(asistencia.sede.nombre))
        ws.cell(row=row_idx, column=4, value=_safe_cell(asistencia.tipo_registro.nombre))
        ws.cell(row=row_idx, column=5, value=asistencia.fecha_registro.isoformat())
        ws.cell(
            row=row_idx,
            column=6,
            value=asistencia.registrado_en.strftime("%H:%M:%S"),
        )

        if novedad is not None:
            ws.cell(row=row_idx, column=7, value=_safe_cell(novedad.tipo_novedad.nombre))
            color_hex = (novedad.tipo_novedad.color or "").lstrip("#")
            if len(color_hex) == 6:
                fill = PatternFill(
                    start_color=color_hex, end_color=color_hex, fill_type="solid"
                )
                font = Font(color="FFFFFF")
                for col in range(1, len(_EXCEL_HEADERS) + 1):
                    ws.cell(row=row_idx, column=col).fill = fill
                    ws.cell(row=row_idx, column=col).font = font

    for col_idx in range(1, len(_EXCEL_HEADERS) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 20

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generar_nombre_archivo(
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
) -> str:
    parts = ["reporte_asistencia"]
    if fecha_desde:
        parts.append(fecha_desde.isoformat())
    if fecha_hasta:
        parts.append(fecha_hasta.isoformat())
    if not fecha_desde and not fecha_hasta:
        parts.append(tz_now().date().isoformat())
    return "_".join(parts) + ".xlsx"
